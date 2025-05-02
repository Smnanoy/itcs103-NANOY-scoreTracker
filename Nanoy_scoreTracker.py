import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os

# File name
filename = "student_scores.xlsx"

# Create workbook if it doesn't exist
if not os.path.exists(filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Scores"
    ws.append(["Name", "Scores", "Remarks"])
    wb.save(filename)

# Load workbook
wb = load_workbook(filename)
ws = wb.active

# Remove previous average
def remove_old_average():
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value == "Average":
            ws.delete_rows(row[0].row)
            break

# Load the latest score
def load_latest_score():
    last_data_row = None
    for row in reversed(list(ws.iter_rows(min_row=2, values_only=True))):
        if row[0] != "Average":
            last_data_row = row
            break

    if last_data_row:
        name, score, _ = last_data_row
        listbox.delete(0, tk.END)
        listbox.insert(tk.END, f"{name} = {score}")

# Save score to Excel
def save_to_excel():
    student_name = name_entry.get().strip()
    score = score_entry.get().strip()
    
    if not student_name or not score:
        messagebox.showerror("Error!", "Both fields are required.")
        return

    try:
        score = float(score)
        if score < 0 or score > 100:
            messagebox.showerror("Error!", "Grade must be between 0 to 100.")
            return
    except ValueError:
        messagebox.showerror("Error!", "Grade must be a number.")
        return

    remarks = "Passed" if score >= 75 else "Failed"
    ws.append([student_name, score, remarks])

    # remove old entry
    remove_old_average()
    # Last row
    last_data_row = 2  # Start after header
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value != "Average":
            last_data_row = row[0].row

    # update average
    ws[f"A{last_data_row + 1}"] = "Average"
    ws[f"B{last_data_row + 1}"] = f"=AVERAGE(B2:B{last_data_row})"
    wb.save(filename)

    # Store data to listbox
    listbox.insert(tk.END, f"{student_name} = {score}")

    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)
    messagebox.showinfo("Finished", "Score recorded successfully!")

# Tkinter Interface
window = tk.Tk()
window.title("Score Tracker")
window.geometry("400x350")
window.configure(background="Gray30")

# Main window frame
main_frame = tk.Frame(window, padx=10, pady=10)
main_frame.pack(expand=True, fill=tk.BOTH)

# Information Frame
input_frame = tk.LabelFrame(main_frame, text="Add New Score", 
                          bd=2, relief="groove", padx=10, pady=10)
input_frame.pack(fill=tk.X, padx=5, pady=5)

# Name entries
tk.Label(input_frame, text="Name:", font=("helvetica", 12)).grid(row=0, column=0, padx=5, pady=5, sticky="e")
name_entry = tk.Entry(input_frame, font=("helvetica", 12))
name_entry.grid(row=0, column=1, padx=5, pady=5, sticky="we")

# Score entries
tk.Label(input_frame, text="Score:", font=("helvetica", 12)).grid(row=1, column=0, padx=5, pady=5, sticky="e")
score_entry = tk.Entry(input_frame, font=("helvetica", 12))
score_entry.grid(row=1, column=1, padx=5, pady=5, sticky="we")

input_frame.grid_columnconfigure(1, weight=1)

# Score button
submit_btn = tk.Button(input_frame, text="Add Score", bg="darkviolet", fg="white", 
                     font=("helvetica", 12, "bold"), command=save_to_excel)
submit_btn.grid(row=2, column=0, columnspan=2, pady=10, sticky="we")

# Score Frame
score_frame = tk.LabelFrame(main_frame, text="Score History", 
                          bd=2, relief="groove", padx=10, pady=10)
score_frame.pack(expand=True, fill=tk.BOTH, padx=5, pady=5)

# Listbox to show scores
listbox = tk.Listbox(score_frame, font=("helvetica", 12))
listbox.pack(expand=True, fill=tk.BOTH, side=tk.LEFT)

# Scrollbar
scrollbar = tk.Scrollbar(score_frame)
scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
listbox.config(yscrollcommand=scrollbar.set)
scrollbar.config(command=listbox.yview)

# Load latest score 
load_latest_score()

window.mainloop()
